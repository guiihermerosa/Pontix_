"""
ExportService — geração de relatórios Excel (XLSX) e CSV.

Funções públicas:
  export_attendance_xlsx()         — relatório geral (todas as marcações + resumo)
  export_employee_xlsx()           — relatório individual de UM funcionário
  export_employees_xlsx()          — cadastro de funcionários
  export_attendance_csv()          — CSV geral de marcações
  export_employees_csv()           — CSV de funcionários

Estrutura do XLSX geral (para proprietário):
  Aba "Capa"        — identidade da empresa + resumo executivo do período
  Aba "Resumo"      — tabela consolidada: funcionário × dias/horas/faltas/atestados
  Aba "Marcações"   — detalhe de cada batida individual
  Aba "Atestados"   — todos os atestados do período

Estrutura do XLSX individual (para contabilidade, 1 por funcionário):
  Aba "Capa"        — dados do funcionário + resumo do período
  Aba "Marcações"   — batidas do funcionário com pares entrada/saída e horas
  Aba "Atestados"   — atestados do funcionário (se houver)
"""

import base64
import calendar
import csv
import io
import logging
from datetime import date, datetime, timedelta
from typing import Optional

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.database.models import AttendanceRecord, Employee, MedicalCertificate, Setting

logger = logging.getLogger("export_service")


# ============================================================================
# PALETA DE CORES
# ============================================================================
_P = {
    "navy":     "0D2137",   # azul naval — header principal
    "blue":     "1A56DB",   # azul — sub-header
    "teal":     "0694A2",   # teal — atestados / destaque
    "green":    "057A55",   # verde — entrada / OK
    "red":      "C81E1E",   # vermelho — saída / falta
    "orange":   "B43403",   # laranja — atenção
    "gray_dk":  "374151",   # cinza escuro — texto
    "gray_md":  "6B7280",   # cinza médio — subtexto
    "gray_lt":  "F3F4F6",   # cinza claro — zebra par
    "blue_lt":  "EBF5FF",   # azul muito claro — zebra atestado
    "yellow":   "FEFCE8",   # amarelo — falta sem justificativa
    "white":    "FFFFFF",
    "border":   "D1D5DB",   # borda
    "total":    "DBEAFE",   # linha de total
    "cert_hdr": "0E7490",   # cabeçalho seção atestados
}


# ============================================================================
# HELPERS BÁSICOS DE ESTILO
# ============================================================================

def _f(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _fo(bold=False, size=10, color="374151", italic=False, name="Calibri") -> Font:
    return Font(bold=bold, size=size, color=color, italic=italic, name=name)


def _b(color="D1D5DB") -> Border:
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _a(h="center", v="center", wrap=False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


_B_MAIN   = _b()
_A_CTR    = _a()
_A_LEFT   = _a(h="left")
_A_RIGHT  = _a(h="right")


def _style(ws, row: int, ncols: int, fill: str, font_kw: dict, align=_A_CTR, height=20):
    fo = _fo(**font_kw)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill      = _f(fill)
        cell.font      = fo
        cell.alignment = align
        cell.border    = _b(_P["white"] if fill != _P["white"] else _P["border"])
    ws.row_dimensions[row].height = height


def _data_row(ws, row: int, ncols: int, alt: bool, height=18):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        if alt:
            cell.fill = _f(_P["gray_lt"])
        cell.font      = _fo(size=10)
        cell.alignment = _A_LEFT
        cell.border    = _B_MAIN
    ws.row_dimensions[row].height = height


def _total_row(ws, row: int, ncols: int, height=20):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill      = _f(_P["total"])
        cell.font      = _fo(bold=True, size=10, color=_P["navy"])
        cell.alignment = _A_CTR
        cell.border    = _b(_P["blue"])
    ws.row_dimensions[row].height = height


def _auto_width(ws, min_w=8, max_w=42):
    for col_cells in ws.columns:
        mx = min_w
        ltr = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                v = len(str(cell.value or ""))
                if v > mx:
                    mx = v
            except Exception:
                pass
        ws.column_dimensions[ltr].width = min(mx + 3, max_w)


def _merge_write(ws, row, c1, c2, value, fill, font_kw, align=_A_LEFT, height=None):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = ws.cell(row=row, column=c1, value=value)
    cell.fill      = _f(fill)
    cell.font      = _fo(**font_kw)
    cell.alignment = align
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).border = _b(_P["white"] if fill != _P["white"] else _P["border"])
    if height:
        ws.row_dimensions[row].height = height


# ============================================================================
# CABEÇALHO CORPORATIVO
# ============================================================================

def _company_header(ws, company: dict, ncols: int, subtitle: str = "") -> int:
    """
    Insere o cabeçalho com identidade visual da empresa.
    Retorna o número da próxima linha livre.
    """
    ws.sheet_view.showGridLines = False

    # Fundo das 5 primeiras linhas
    for r in range(1, 6):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill   = _f(_P["navy"])
            ws.cell(row=r, column=c).border = _b(_P["white"])

    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 5   # separador

    # Logo
    logo_offset = 0
    logo_b64 = company.get("logo_b64", "")
    if logo_b64:
        try:
            raw = logo_b64.split(",")[-1]
            xl = XLImage(io.BytesIO(base64.b64decode(raw)))
            xl.width, xl.height = 140, 56
            ws.add_image(xl, "A1")
            logo_offset = 3
        except Exception as exc:
            logger.warning("Logo XLSX: %s", exc)

    tc = logo_offset + 1

    def _hc(r, c, v, bold=False, sz=11, color=_P["white"], italic=False):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font      = _fo(bold=bold, size=sz, color=color, italic=italic)
        cell.alignment = _A_LEFT

    name    = company.get("name") or "Empresa"
    cnpj    = company.get("cnpj") or ""
    address = company.get("address") or ""
    gen     = datetime.now().strftime("%d/%m/%Y às %H:%M")

    _hc(1, tc, name, bold=True, sz=18)
    info = "  ·  ".join(filter(None, [f"CNPJ: {cnpj}" if cnpj else "", address]))
    _hc(2, tc, info, sz=9, italic=True, color="93C5FD")
    _hc(3, tc, subtitle or "Relatório de Ponto", bold=True, sz=13)
    _hc(4, tc, f"Gerado em: {gen}", sz=8, italic=True, color="6B7280")

    # Linha separadora azul médio
    for c in range(1, ncols + 1):
        ws.cell(row=5, column=c).fill = _f(_P["blue"])
    ws.row_dimensions[5].height = 3

    return 6


# ============================================================================
# LEITURA DE DADOS
# ============================================================================

def _load_company(db: Session) -> dict:
    def _g(key, default=""):
        row = db.query(Setting).filter(Setting.key == key).first()
        return row.value if row and row.value else default

    return {
        "name":    _g("company_name"),
        "cnpj":    _g("company_cnpj"),
        "address": _g("company_address"),
        "logo_b64": _g("company_logo_b64"),
    }


def _query_records(db, start_date=None, end_date=None, userid=None, record_type=None):
    q = db.query(AttendanceRecord)
    if start_date:  q = q.filter(AttendanceRecord.attendance_date >= start_date)
    if end_date:    q = q.filter(AttendanceRecord.attendance_date <= end_date)
    if userid:      q = q.filter(AttendanceRecord.userid == userid)
    if record_type: q = q.filter(AttendanceRecord.record_type == record_type)
    return q.order_by(
        AttendanceRecord.attendance_date.asc(),
        AttendanceRecord.attendance_time.asc(),
    ).all()


def _query_certs(db, userid=None, start_date=None, end_date=None):
    q = db.query(MedicalCertificate)
    if userid:     q = q.filter(MedicalCertificate.userid == userid)
    if start_date: q = q.filter(MedicalCertificate.cert_date >= start_date)
    if end_date:   q = q.filter(MedicalCertificate.cert_date <= end_date)
    return q.order_by(MedicalCertificate.cert_date.asc()).all()


def _date_range(start_date: str, end_date: str) -> list[str]:
    """Retorna lista de datas (strings YYYY-MM-DD) entre start e end."""
    try:
        d0 = datetime.strptime(start_date, "%Y-%m-%d").date()
        d1 = datetime.strptime(end_date,   "%Y-%m-%d").date()
        return [(d0 + timedelta(days=i)).isoformat() for i in range((d1 - d0).days + 1)]
    except Exception:
        return []


def _is_weekday(date_str: str) -> bool:
    """Retorna True se a data não for sábado (5) nem domingo (6)."""
    try:
        return datetime.strptime(date_str, "%Y-%m-%d").weekday() < 5
    except Exception:
        return True


def _calc_hours(entry: str, exit_: str) -> float:
    """Calcula horas trabalhadas entre duas strings HH:MM:SS."""
    try:
        fmt = "%H:%M:%S"
        t0 = datetime.strptime(entry, fmt)
        t1 = datetime.strptime(exit_, fmt)
        delta = t1 - t0
        if delta.total_seconds() < 0:
            return 0.0
        return round(delta.total_seconds() / 3600, 2)
    except Exception:
        return 0.0


def _build_daily_summary(records: list, certs: list, all_dates: list) -> list[dict]:
    """
    Constrói resumo diário para UM funcionário.
    Retorna lista de dicts por data com: entry, exit, hours, absent, justified.
    """
    # Agrupa batidas por dia
    by_day: dict[str, list] = {}
    for r in records:
        by_day.setdefault(r.attendance_date, []).append(r)

    # Datas de atestado
    cert_dates: dict[str, MedicalCertificate] = {}
    for cert in certs:
        d0 = datetime.strptime(cert.cert_date, "%Y-%m-%d").date()
        for i in range(cert.days_off):
            d = (d0 + timedelta(days=i)).isoformat()
            cert_dates[d] = cert

    rows = []
    for d in all_dates:
        if not _is_weekday(d):
            continue
        day_recs = sorted(by_day.get(d, []), key=lambda r: r.attendance_time)
        entry = exit_ = ""
        hours = 0.0

        if day_recs:
            # Primeiro ENTRY e último EXIT
            entries = [r for r in day_recs if r.record_type == "ENTRY"]
            exits   = [r for r in day_recs if r.record_type == "EXIT"]
            if entries: entry  = entries[0].attendance_time
            if exits:   exit_  = exits[-1].attendance_time
            if entry and exit_: hours = _calc_hours(entry, exit_)

        absent     = not bool(day_recs)
        justified  = d in cert_dates
        cert       = cert_dates.get(d)

        rows.append({
            "date":      d,
            "entry":     entry,
            "exit":      exit_,
            "hours":     hours,
            "absent":    absent,
            "justified": justified,
            "cert_obs":  cert.description if cert else "",
        })
    return rows


# ============================================================================
# ABA DE MARCAÇÕES — uso individual ou geral
# ============================================================================

def _sheet_marks(ws, records, company, subtitle, emp_map, certs_by_userid,
                 all_dates_or_none=None):
    """
    Aba de marcações detalhadas.
    Se all_dates_or_none fornecido, mostra também dias sem batida (faltas).
    """
    ncols = 9
    ws.sheet_properties.tabColor = _P["blue"]
    company["subtitle"] = subtitle
    nr = _company_header(ws, company, ncols)

    # Cabeçalho de colunas
    headers = ["Data", "Dia", "Funcionário", "Entrada", "Saída",
               "Horas", "Situação", "Atestado", "Obs. Atestado"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=nr, column=c, value=h)
    _style(ws, nr, ncols, _P["navy"], {"bold": True, "color": _P["white"], "size": 10}, height=24)
    nr += 1

    # Datas de atestados
    cert_dates: dict[str, str] = {}
    for uid, certs in certs_by_userid.items():
        for cert in certs:
            d0 = datetime.strptime(cert.cert_date, "%Y-%m-%d").date()
            for i in range(cert.days_off):
                d = (d0 + timedelta(days=i)).isoformat()
                cert_dates[d] = cert.description or "Atestado"

    if all_dates_or_none:
        # Modo individual: uma linha por dia do período
        userid = records[0].userid if records else ""
        emp    = emp_map.get(userid)
        nome   = emp.name if emp else userid
        daily  = _build_daily_summary(records, certs_by_userid.get(userid, []), all_dates_or_none)
        idx = 0
        for day in daily:
            d     = day["date"]
            wkday = ["Seg","Ter","Qua","Qui","Sex","Sáb","Dom"][datetime.strptime(d,"%Y-%m-%d").weekday()]
            entry = day["entry"][:5] if day["entry"] else ""
            exit_ = day["exit"][:5]  if day["exit"]  else ""
            hours_s = f"{day['hours']:.1f}h" if day["hours"] else ""

            if day["absent"] and day["justified"]:
                sit_txt  = "Atestado"
                sit_fill = _P["teal"]
                sit_col  = _P["white"]
            elif day["absent"]:
                sit_txt  = "Falta"
                sit_fill = _P["red"]
                sit_col  = _P["white"]
            elif not entry:
                sit_txt  = "Incompleto"
                sit_fill = _P["orange"]
                sit_col  = _P["white"]
            else:
                sit_txt  = "Presente"
                sit_fill = _P["green"]
                sit_col  = _P["white"]

            row_data = [d, wkday, nome, entry, exit_, hours_s, sit_txt,
                        "Sim" if day["justified"] else "", day["cert_obs"]]
            alt = (idx % 2 == 0)
            for c, val in enumerate(row_data, 1):
                cell = ws.cell(row=nr, column=c, value=val)
                cell.font      = _fo(size=10)
                cell.alignment = _A_CTR if c not in (3, 9) else _A_LEFT
                cell.border    = _B_MAIN
                if alt:
                    cell.fill  = _f(_P["gray_lt"])
                # Coluna Situação com cor própria
                if c == 7:
                    cell.fill  = _f(sit_fill)
                    cell.font  = _fo(bold=True, size=10, color=sit_col)
            ws.row_dimensions[nr].height = 18
            nr += 1
            idx += 1
    else:
        # Modo geral: todos os registros
        for idx, rec in enumerate(records):
            d     = rec.attendance_date
            wkday = ["Seg","Ter","Qua","Qui","Sex","Sáb","Dom"][datetime.strptime(d,"%Y-%m-%d").weekday()]
            emp   = emp_map.get(rec.userid)
            nome  = rec.employee_name or (emp.name if emp else rec.userid)
            tipo  = "Entrada" if rec.record_type == "ENTRY" else "Saída" if rec.record_type == "EXIT" else "—"
            entry = rec.attendance_time[:5] if rec.record_type == "ENTRY" else ""
            exit_ = rec.attendance_time[:5] if rec.record_type == "EXIT"  else ""
            has_cert = d in cert_dates
            cert_obs = cert_dates.get(d, "")

            row_data = [d, wkday, nome, entry, exit_, "", tipo, "Sim" if has_cert else "", cert_obs]
            alt = (idx % 2 == 0)
            for c, val in enumerate(row_data, 1):
                cell = ws.cell(row=nr, column=c, value=val)
                cell.font      = _fo(size=10)
                cell.alignment = _A_CTR if c not in (3, 9) else _A_LEFT
                cell.border    = _B_MAIN
                if alt: cell.fill = _f(_P["gray_lt"])
                if c == 7:
                    if rec.record_type == "ENTRY":
                        cell.fill = _f(_P["green"]); cell.font = _fo(bold=True, color=_P["white"])
                    elif rec.record_type == "EXIT":
                        cell.fill = _f(_P["red"]);   cell.font = _fo(bold=True, color=_P["white"])
            ws.row_dimensions[nr].height = 18
            nr += 1

    # Rodapé
    total_txt = f"Total: {len(records)} marcações"
    for c in range(1, ncols + 1):
        ws.cell(row=nr, column=c).fill = _f(_P["total"])
        ws.cell(row=nr, column=c).border = _b(_P["blue"])
    ws.cell(row=nr, column=1, value=total_txt).font = _fo(bold=True, color=_P["navy"])
    ws.row_dimensions[nr].height = 20

    _auto_width(ws)
    ws.freeze_panes = "A7"


# ============================================================================
# ABA DE ATESTADOS
# ============================================================================

def _sheet_certs(ws, certs, company, subtitle, emp_map):
    ncols = 7
    ws.sheet_properties.tabColor = _P["teal"]
    company["subtitle"] = subtitle
    nr = _company_header(ws, company, ncols)

    headers = ["Data", "Funcionário", "ID", "Dias Afastamento", "Observação / Diagnóstico",
               "Arquivo", "Cadastrado em"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=nr, column=c, value=h)
    _style(ws, nr, ncols, _P["cert_hdr"], {"bold": True, "color": _P["white"], "size": 10}, height=24)
    nr += 1

    if not certs:
        ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=ncols)
        cell = ws.cell(row=nr, column=1, value="Nenhum atestado registrado no período.")
        cell.font = _fo(italic=True, color=_P["gray_md"])
        cell.alignment = _A_CTR
        ws.row_dimensions[nr].height = 24
        return

    for idx, cert in enumerate(certs):
        emp  = emp_map.get(cert.userid)
        nome = emp.name if emp else cert.userid
        cat  = cert.created_at.strftime("%d/%m/%Y") if cert.created_at else ""
        has_file = "✓ Anexado" if cert.file_b64 else "—"

        row_data = [cert.cert_date, nome, cert.userid, cert.days_off,
                    cert.description or "—", has_file, cat]
        alt = (idx % 2 == 0)
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(row=nr, column=c, value=val)
            cell.font      = _fo(size=10)
            cell.alignment = _A_CTR if c not in (5,) else _A_LEFT
            cell.border    = _B_MAIN
            if alt: cell.fill = _f(_P["blue_lt"])
            if c == 6 and cert.file_b64:
                cell.font = _fo(bold=True, color=_P["teal"])
        ws.row_dimensions[nr].height = 18
        nr += 1

    _total_row(ws, nr, ncols)
    ws.cell(row=nr, column=1, value=f"Total: {len(certs)} atestado(s)").font = _fo(bold=True, color=_P["navy"])
    ws.cell(row=nr, column=4, value=sum(c.days_off for c in certs)).font = _fo(bold=True, color=_P["navy"])
    ws.cell(row=nr, column=4).alignment = _A_CTR

    _auto_width(ws)
    ws.freeze_panes = "A7"


# ============================================================================
# ABA DE RESUMO GERAL
# ============================================================================

def _sheet_summary(ws, records, certs, company, start_date, end_date, emp_map):
    ncols = 10
    ws.sheet_properties.tabColor = _P["teal"]
    company["subtitle"] = f"Resumo Consolidado  ·  {start_date or 'Início'} → {end_date or 'Hoje'}"
    nr = _company_header(ws, company, ncols)

    headers = ["Funcionário", "ID", "Dias Úteis", "Dias Presentes",
               "Faltas", "Faltas Justif.", "Entradas", "Saídas",
               "Horas Trab. (est.)", "Atestados"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=nr, column=c, value=h)
    _style(ws, nr, ncols, _P["navy"], {"bold": True, "color": _P["white"], "size": 10}, height=26)
    nr += 1

    all_dates = _date_range(start_date or date.today().replace(day=1).isoformat(),
                            end_date   or date.today().isoformat())
    work_days = sum(1 for d in all_dates if _is_weekday(d))

    # Agrupa por userid
    from collections import defaultdict
    by_uid: dict[str, list] = defaultdict(list)
    for r in records:
        by_uid[r.userid].append(r)

    cert_by_uid: dict[str, list] = defaultdict(list)
    for c in certs:
        cert_by_uid[c.userid].append(c)

    totais = {"present": 0, "absent": 0, "justified": 0, "entries": 0, "exits": 0, "hours": 0.0, "certs": 0}

    uids = sorted(by_uid.keys(), key=lambda u: (emp_map[u].name if u in emp_map else u))
    for idx, uid in enumerate(uids):
        emp     = emp_map.get(uid)
        nome    = emp.name if emp else uid
        recs    = by_uid[uid]
        c_list  = cert_by_uid.get(uid, [])

        # Datas de atestado para este funcionário
        cert_dates_set: set[str] = set()
        for cert in c_list:
            d0 = datetime.strptime(cert.cert_date, "%Y-%m-%d").date()
            for i in range(cert.days_off):
                cert_dates_set.add((d0 + timedelta(days=i)).isoformat())

        days_present  = len({r.attendance_date for r in recs})
        entries_count = sum(1 for r in recs if r.record_type == "ENTRY")
        exits_count   = sum(1 for r in recs if r.record_type == "EXIT")

        # Horas estimadas por pares entrada/saída
        total_hours = 0.0
        by_day_recs: dict[str, list] = defaultdict(list)
        for r in recs:
            by_day_recs[r.attendance_date].append(r)
        for day_recs in by_day_recs.values():
            day_entries = sorted([r for r in day_recs if r.record_type == "ENTRY"], key=lambda x: x.attendance_time)
            day_exits   = sorted([r for r in day_recs if r.record_type == "EXIT"],  key=lambda x: x.attendance_time)
            for i in range(min(len(day_entries), len(day_exits))):
                total_hours += _calc_hours(day_entries[i].attendance_time, day_exits[i].attendance_time)

        absent_days    = max(0, work_days - days_present)
        justified_days = len([d for d in cert_dates_set if _is_weekday(d)])
        unjust_days    = max(0, absent_days - justified_days)

        row_data = [
            nome, uid, work_days, days_present,
            absent_days, justified_days,
            entries_count, exits_count,
            f"{total_hours:.1f}h", len(c_list),
        ]
        alt = (idx % 2 == 0)
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(row=nr, column=c, value=val)
            cell.alignment = _A_LEFT if c <= 2 else _A_CTR
            cell.border    = _B_MAIN
            cell.font      = _fo(size=10)
            if alt: cell.fill = _f(_P["gray_lt"])
            # Faltas sem justificativa em vermelho
            if c == 5 and unjust_days > 0:
                cell.fill = _f(_P["red"]); cell.font = _fo(bold=True, color=_P["white"])
            if c == 6 and justified_days > 0:
                cell.fill = _f(_P["teal"]); cell.font = _fo(bold=True, color=_P["white"])
        ws.row_dimensions[nr].height = 20

        totais["present"]   += days_present
        totais["absent"]    += absent_days
        totais["justified"] += justified_days
        totais["entries"]   += entries_count
        totais["exits"]     += exits_count
        totais["hours"]     += total_hours
        totais["certs"]     += len(c_list)
        nr += 1

    # Linha de total
    _total_row(ws, nr, ncols)
    tot_vals = ["TOTAL", f"{len(uids)} func.", work_days * len(uids),
                totais["present"], totais["absent"], totais["justified"],
                totais["entries"], totais["exits"],
                f"{totais['hours']:.1f}h", totais["certs"]]
    for c, val in enumerate(tot_vals, 1):
        cell = ws.cell(row=nr, column=c, value=val)
        cell.font      = _fo(bold=True, color=_P["navy"])
        cell.alignment = _A_LEFT if c <= 2 else _A_CTR

    _auto_width(ws)
    ws.freeze_panes = "A7"


# ============================================================================
# ABA CAPA — executivo
# ============================================================================

def _sheet_cover(ws, company: dict, title: str, info_lines: list[tuple]):
    """
    Capa executiva com identidade visual e caixas de informação.
    info_lines: lista de (label, valor) para exibir.
    """
    ncols = 6
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _P["navy"]

    # Fundo total navy
    for r in range(1, 20):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill   = _f(_P["navy"])
            ws.cell(row=r, column=c).border = _b(_P["navy"])
        ws.row_dimensions[r].height = 22

    ws.row_dimensions[1].height = 14
    ws.row_dimensions[2].height = 50

    # Logo
    logo_b64 = company.get("logo_b64", "")
    if logo_b64:
        try:
            raw = logo_b64.split(",")[-1]
            xl = XLImage(io.BytesIO(base64.b64decode(raw)))
            xl.width, xl.height = 160, 64
            ws.add_image(xl, "B2")
        except Exception:
            pass

    # Nome da empresa
    ws.row_dimensions[3].height = 14
    ws.row_dimensions[4].height = 40
    ws.merge_cells("B4:F4")
    cell = ws.cell(row=4, column=2, value=company.get("name") or "Empresa")
    cell.font      = _fo(bold=True, size=22, color=_P["white"])
    cell.alignment = _A_LEFT

    cnpj = company.get("cnpj", "")
    if cnpj:
        ws.row_dimensions[5].height = 20
        ws.merge_cells("B5:F5")
        cell = ws.cell(row=5, column=2, value=f"CNPJ: {cnpj}")
        cell.font = _fo(size=11, color="93C5FD", italic=True)
        cell.alignment = _A_LEFT

    # Título do relatório
    ws.row_dimensions[7].height = 36
    ws.merge_cells("B7:F7")
    cell = ws.cell(row=7, column=2, value=title)
    cell.font      = _fo(bold=True, size=16, color="BFDBFE")
    cell.alignment = _A_LEFT

    # Linha divisória
    ws.row_dimensions[8].height = 3
    for c in range(2, 7):
        ws.cell(row=8, column=c).fill = _f(_P["blue"])

    # Caixas de informação
    row = 10
    for label, value in info_lines:
        ws.row_dimensions[row].height = 28
        # Label
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        lc = ws.cell(row=row, column=2, value=label)
        lc.font      = _fo(size=10, color="93C5FD", italic=True)
        lc.alignment = _A_RIGHT
        # Valor
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        vc = ws.cell(row=row, column=4, value=value)
        vc.font      = _fo(bold=True, size=11, color=_P["white"])
        vc.alignment = _A_LEFT
        row += 1

    # Footer
    for r in range(row + 1, row + 4):
        ws.row_dimensions[r].height = 22
    ws.merge_cells(start_row=row + 2, start_column=2, end_row=row + 2, end_column=6)
    fc = ws.cell(row=row + 2, column=2,
                 value=f"Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')} · Pontix · GRB Tecnologia")
    fc.font      = _fo(size=9, color="4B5563", italic=True)
    fc.alignment = _A_CTR

    for c in range(1, 7):
        ws.column_dimensions[get_column_letter(c)].width = 18


# ============================================================================
# EXPORT GERAL (para o proprietário)
# ============================================================================

def export_attendance_xlsx(
    db: Session,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    userid: Optional[str] = None,
    record_type: Optional[str] = None,
) -> bytes:
    """
    XLSX geral — 4 abas: Capa, Resumo, Marcações, Atestados.
    Destinado ao proprietário da empresa.
    """
    records = _query_records(db, start_date, end_date, userid, record_type)
    certs   = _query_certs(db, userid, start_date, end_date)
    company = _load_company(db)
    emp_map = {e.userid: e for e in db.query(Employee).all()}

    period  = f"{start_date or 'Início'} → {end_date or 'Hoje'}"
    nb_emp  = len({r.userid for r in records})
    nb_cert = len(certs)

    wb = Workbook()

    # ── Capa ─────────────────────────────────────────────────────────────
    ws_cov = wb.active
    ws_cov.title = "Capa"
    _sheet_cover(ws_cov, company, "Relatório Geral de Ponto", [
        ("Período",        period),
        ("Funcionários",   str(nb_emp)),
        ("Total Batidas",  str(len(records))),
        ("Atestados",      str(nb_cert)),
        ("Gerado em",      datetime.now().strftime("%d/%m/%Y às %H:%M")),
    ])

    # ── Resumo ────────────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("Resumo")
    _sheet_summary(ws_sum, records, certs, company, start_date, end_date, emp_map)

    # ── Marcações ─────────────────────────────────────────────────────────
    ws_det = wb.create_sheet("Marcações")
    certs_by_uid = {}
    for c in certs:
        certs_by_uid.setdefault(c.userid, []).append(c)
    _sheet_marks(ws_det, records, company,
                 f"Marcações  ·  {period}", emp_map, certs_by_uid)

    # ── Atestados ─────────────────────────────────────────────────────────
    ws_cert = wb.create_sheet("Atestados")
    _sheet_certs(ws_cert, certs, company,
                 f"Atestados Médicos  ·  {period}", emp_map)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    logger.info("XLSX geral: %d marcações, %d atestados", len(records), len(certs))
    return buf.getvalue()


# ============================================================================
# EXPORT INDIVIDUAL (para contabilidade)
# ============================================================================

def export_employee_xlsx(
    db: Session,
    userid: str,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
) -> bytes:
    """
    XLSX individual de UM funcionário — 3 abas: Capa, Marcações, Atestados.
    Destinado ao escritório de contabilidade.
    """
    records = _query_records(db, start_date, end_date, userid)
    certs   = _query_certs(db, userid, start_date, end_date)
    company = _load_company(db)
    emp_map = {e.userid: e for e in db.query(Employee).all()}
    emp     = emp_map.get(userid)
    nome    = emp.name if emp else userid

    period     = f"{start_date or 'Início'} → {end_date or 'Hoje'}"
    all_dates  = _date_range(
        start_date or date.today().replace(day=1).isoformat(),
        end_date   or date.today().isoformat(),
    )
    work_days  = sum(1 for d in all_dates if _is_weekday(d))
    days_pres  = len({r.attendance_date for r in records})
    total_h    = 0.0
    from collections import defaultdict
    by_day: dict[str, list] = defaultdict(list)
    for r in records:
        by_day[r.attendance_date].append(r)
    for day_recs in by_day.values():
        entr = sorted([r for r in day_recs if r.record_type == "ENTRY"], key=lambda x: x.attendance_time)
        exts = sorted([r for r in day_recs if r.record_type == "EXIT"],  key=lambda x: x.attendance_time)
        for i in range(min(len(entr), len(exts))):
            total_h += _calc_hours(entr[i].attendance_time, exts[i].attendance_time)

    cert_days = set()
    for cert in certs:
        d0 = datetime.strptime(cert.cert_date, "%Y-%m-%d").date()
        for i in range(cert.days_off):
            cert_days.add((d0 + timedelta(days=i)).isoformat())
    absent    = max(0, work_days - days_pres)
    justified = len([d for d in cert_days if _is_weekday(d)])

    wb = Workbook()

    # ── Capa ─────────────────────────────────────────────────────────────
    ws_cov = wb.active
    ws_cov.title = "Capa"
    dept_s = str(emp.department) if emp else "—"
    _sheet_cover(ws_cov, company, f"Relatório de Ponto — {nome}", [
        ("Funcionário",     nome),
        ("ID / Matrícula",  userid),
        ("Departamento",    dept_s),
        ("Período",         period),
        ("Dias úteis",      str(work_days)),
        ("Dias presentes",  str(days_pres)),
        ("Faltas",          str(absent)),
        ("Justificadas",    str(justified)),
        ("Horas (est.)",    f"{total_h:.1f}h"),
        ("Atestados",       str(len(certs))),
    ])

    # ── Marcações ─────────────────────────────────────────────────────────
    ws_det = wb.create_sheet("Marcações")
    certs_by_uid = {userid: certs}
    _sheet_marks(ws_det, records, company,
                 f"Marcações — {nome}  ·  {period}",
                 emp_map, certs_by_uid, all_dates_or_none=all_dates)

    # ── Atestados ─────────────────────────────────────────────────────────
    ws_cert = wb.create_sheet("Atestados")
    _sheet_certs(ws_cert, certs, company,
                 f"Atestados — {nome}  ·  {period}", emp_map)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    logger.info("XLSX individual gerado: userid=%s, %d marcações", userid, len(records))
    return buf.getvalue()


# ============================================================================
# EXPORT CADASTRO DE FUNCIONÁRIOS
# ============================================================================

def export_employees_xlsx(db: Session) -> bytes:
    employees = db.query(Employee).order_by(Employee.name).all()
    company   = _load_company(db)

    wb  = Workbook()
    ws  = wb.active
    ws.title = "Funcionários"
    ncols = 10
    ws.sheet_properties.tabColor = _P["navy"]

    company["subtitle"] = f"Cadastro de Funcionários  ·  Total: {len(employees)}"
    nr = _company_header(ws, company, ncols)

    headers = ["ID", "Matrícula", "Nome", "Departamento", "Cargo",
               "Turno", "Cartão", "Face", "Dig./Palma", "Status Sync"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=nr, column=c, value=h)
    _style(ws, nr, ncols, _P["navy"], {"bold": True, "color": _P["white"], "size": 10}, height=24)
    nr += 1

    sync_colors = {"SUCCESS": _P["green"], "ERROR": _P["red"], "PENDING": _P["orange"]}

    for idx, emp in enumerate(employees):
        row_data = [
            emp.id, emp.userid, emp.name or "",
            emp.department, _role_label(emp.role), emp.schedule,
            emp.access_card_number or "",
            "✓" if emp.has_face else "—",
            "✓" if (emp.has_fingerprint or emp.has_palm) else "—",
            _sync_label(emp.sync_status),
        ]
        alt = (idx % 2 == 0)
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(row=nr, column=c, value=val)
            cell.font      = _fo(size=10)
            cell.alignment = _A_LEFT if c == 3 else _A_CTR
            cell.border    = _B_MAIN
            if alt: cell.fill = _f(_P["gray_lt"])
            if c == 10:
                col = sync_colors.get(emp.sync_status or "", _P["gray_md"])
                cell.font = _fo(bold=True, color=col)
        ws.row_dimensions[nr].height = 18
        nr += 1

    _total_row(ws, nr, ncols)
    ws.cell(row=nr, column=1, value=f"Total: {len(employees)} funcionários"
            ).font = _fo(bold=True, color=_P["navy"])
    _auto_width(ws)
    ws.freeze_panes = "A7"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    logger.info("XLSX funcionários: %d registros", len(employees))
    return buf.getvalue()


# ============================================================================
# CSV (mantidos para compatibilidade)
# ============================================================================

def export_attendance_csv(
    db: Session,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    userid: Optional[str] = None,
    record_type: Optional[str] = None,
) -> bytes:
    records = _query_records(db, start_date, end_date, userid, record_type)
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["ID", "Funcionário", "Nome", "Data", "Hora", "Tipo", "Origem"])
    for rec in records:
        w.writerow([rec.id, rec.userid, rec.employee_name or "",
                    rec.attendance_date, rec.attendance_time,
                    _type_label(rec.record_type), rec.source])
    return out.getvalue().encode("utf-8-sig")


def export_employees_csv(db: Session) -> bytes:
    employees = db.query(Employee).order_by(Employee.name).all()
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["ID", "Matrícula", "Nome", "Departamento", "Cargo",
                "Turno", "Cartão", "Face", "Status Sync", "Última Sync"])
    for emp in employees:
        w.writerow([
            emp.id, emp.userid, emp.name or "",
            emp.department, _role_label(emp.role), emp.schedule,
            emp.access_card_number or "",
            "Sim" if emp.has_face else "Não",
            _sync_label(emp.sync_status),
            emp.last_synced_at.strftime("%d/%m/%Y %H:%M") if emp.last_synced_at else "",
        ])
    return out.getvalue().encode("utf-8-sig")


# ============================================================================
# LABELS
# ============================================================================

def _type_label(t: str) -> str:
    return {"ENTRY": "Entrada", "EXIT": "Saída", "UNKNOWN": "Desconhecido"}.get(t, t)

def _role_label(role: int) -> str:
    return {0: "Usuário", 1: "Administrador"}.get(role, str(role))

def _sync_label(status: str) -> str:
    return {"PENDING": "Pendente", "SUCCESS": "Sincronizado",
            "ERROR": "Erro", "PROCESSING": "Processando"}.get(status or "", status or "—")
